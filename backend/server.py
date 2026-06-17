"""FlexiDoc backend — universal document conversion API."""
from __future__ import annotations

import io
import logging
import os
import shutil
import subprocess
import tempfile
import time
import uuid
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import List

from dotenv import load_dotenv
from fastapi import APIRouter, BackgroundTasks, FastAPI, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse
from motor.motor_asyncio import AsyncIOMotorClient
from pydantic import BaseModel
from starlette.middleware.cors import CORSMiddleware

import img2pdf
import pikepdf
import fitz  # PyMuPDF
from PIL import Image
from pdf2docx import Converter as PdfToDocxConverter
from pdf2image import convert_from_path
from pypdf import PdfReader, PdfWriter
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / ".env")

mongo_url = os.environ["MONGO_URL"]
mongo_client = AsyncIOMotorClient(mongo_url)
db = mongo_client[os.environ["DB_NAME"]]

STORAGE_DIR = Path(tempfile.gettempdir()) / "flexidoc"
STORAGE_DIR.mkdir(parents=True, exist_ok=True)
RETENTION_SECONDS = 30 * 60

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
logger = logging.getLogger("flexidoc")

app = FastAPI(title="FlexiDoc")
api = APIRouter(prefix="/api")


class ConvertResponse(BaseModel):
    file_id: str
    filename: str
    size: int
    expires_at: str


class ToolInfo(BaseModel):
    id: str
    title: str
    description: str
    category: str
    accept: str
    multiple: bool = False
    icon: str


TOOLS: List[ToolInfo] = [
    ToolInfo(id="pdf-to-word", title="PDF to Word", description="Convert PDF documents into editable DOCX files.", category="convert-from-pdf", accept=".pdf", icon="FileText"),
    ToolInfo(id="word-to-pdf", title="Word to PDF", description="Turn DOC or DOCX files into clean PDFs.", category="convert-to-pdf", accept=".doc,.docx", icon="FileType"),
    ToolInfo(id="pdf-to-jpg", title="PDF to JPG", description="Extract every page of a PDF as a JPG image.", category="convert-from-pdf", accept=".pdf", icon="Image"),
    ToolInfo(id="jpg-to-pdf", title="JPG to PDF", description="Combine images into a single PDF document.", category="convert-to-pdf", accept=".jpg,.jpeg,.png,.webp,.bmp", multiple=True, icon="ImagePlus"),
    ToolInfo(id="merge-pdf", title="Merge PDF", description="Combine multiple PDFs into one ordered document.", category="organize", accept=".pdf", multiple=True, icon="Combine"),
    ToolInfo(id="split-pdf", title="Split PDF", description="Split a PDF into separate single-page files.", category="organize", accept=".pdf", icon="Split"),
    ToolInfo(id="compress-pdf", title="Compress PDF", description="Reduce PDF file size while keeping quality.", category="optimize", accept=".pdf", icon="Minimize2"),
    ToolInfo(id="pdf-to-excel", title="PDF to Excel", description="Pull tables and text from PDF into XLSX sheets.", category="convert-from-pdf", accept=".pdf", icon="Sheet"),
    ToolInfo(id="excel-to-pdf", title="Excel to PDF", description="Convert XLS or XLSX spreadsheets into PDFs.", category="convert-to-pdf", accept=".xls,.xlsx", icon="FileSpreadsheet"),
    ToolInfo(id="pdf-to-powerpoint", title="PDF to PowerPoint", description="Turn each PDF page into a slide deck.", category="convert-from-pdf", accept=".pdf", icon="Presentation"),
    ToolInfo(id="powerpoint-to-pdf", title="PowerPoint to PDF", description="Convert PPT or PPTX decks to shareable PDFs.", category="convert-to-pdf", accept=".ppt,.pptx", icon="MonitorPlay"),
    ToolInfo(id="image-converter", title="Image Converter", description="Convert images between PNG, JPG, WEBP and BMP.", category="image", accept=".jpg,.jpeg,.png,.webp,.bmp", icon="Images"),
]


def _new_job_dir() -> tuple[str, Path]:
    job_id = uuid.uuid4().hex
    d = STORAGE_DIR / job_id
    d.mkdir(parents=True, exist_ok=True)
    return job_id, d


def _save_upload(upload: UploadFile, dest_dir: Path) -> Path:
    safe_name = Path(upload.filename or "upload").name
    path = dest_dir / safe_name
    with path.open("wb") as f:
        shutil.copyfileobj(upload.file, f)
    return path


def _build_response(out_file: Path) -> ConvertResponse:
    file_id = out_file.parent.name
    expires_at = datetime.fromtimestamp(time.time() + RETENTION_SECONDS, tz=timezone.utc).isoformat()
    return ConvertResponse(file_id=file_id, filename=out_file.name, size=out_file.stat().st_size, expires_at=expires_at)


def _cleanup_old_jobs() -> None:
    now = time.time()
    for d in STORAGE_DIR.iterdir():
        if d.is_dir() and now - d.stat().st_mtime > RETENTION_SECONDS:
            shutil.rmtree(d, ignore_errors=True)


def _run_libreoffice(src: Path, out_dir: Path, target_format: str = "pdf") -> Path:
    cmd = ["soffice", "--headless", "--nologo", "--nofirststartwizard",
           "--convert-to", target_format, "--outdir", str(out_dir), str(src)]
    proc = subprocess.run(cmd, capture_output=True, text=True, timeout=180)
    if proc.returncode != 0:
        raise HTTPException(status_code=500, detail="Office conversion failed")
    base = src.stem
    candidates = [c for c in out_dir.glob(f"{base}.*") if c.suffix.lower().lstrip(".") == target_format.split(":")[0]]
    if not candidates:
        raise HTTPException(status_code=500, detail="Converted file not found")
    return candidates[0]


@api.get("/")
async def root():
    return {"app": "FlexiDoc", "status": "ok"}


@api.get("/tools", response_model=List[ToolInfo])
async def list_tools():
    return TOOLS


@api.get("/tools/{tool_id}", response_model=ToolInfo)
async def get_tool(tool_id: str):
    for t in TOOLS:
        if t.id == tool_id:
            return t
    raise HTTPException(status_code=404, detail="Tool not found")


@api.post("/convert/pdf-to-word", response_model=ConvertResponse)
async def pdf_to_word(file: UploadFile = File(...)):
    job_id, d = _new_job_dir()
    src = _save_upload(file, d)
    out = d / f"{src.stem}.docx"
    try:
        cv = PdfToDocxConverter(str(src))
        cv.convert(str(out))
        cv.close()
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Conversion failed: {e}")
    return _build_response(out)


@api.post("/convert/word-to-pdf", response_model=ConvertResponse)
async def word_to_pdf(file: UploadFile = File(...)):
    job_id, d = _new_job_dir()
    src = _save_upload(file, d)
    out = _run_libreoffice(src, d, "pdf")
    return _build_response(out)


@api.post("/convert/pdf-to-jpg", response_model=ConvertResponse)
async def pdf_to_jpg(file: UploadFile = File(...)):
    job_id, d = _new_job_dir()
    src = _save_upload(file, d)
    images = convert_from_path(str(src), dpi=150)
    if not images:
        raise HTTPException(status_code=400, detail="No pages found in PDF")
    if len(images) == 1:
        out = d / f"{src.stem}.jpg"
        images[0].save(out, "JPEG", quality=90)
    else:
        out = d / f"{src.stem}_pages.zip"
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf:
            for i, img in enumerate(images, 1):
                buf = io.BytesIO()
                img.save(buf, "JPEG", quality=90)
                zf.writestr(f"{src.stem}_page_{i}.jpg", buf.getvalue())
    return _build_response(out)


@api.post("/convert/jpg-to-pdf", response_model=ConvertResponse)
async def jpg_to_pdf(files: List[UploadFile] = File(...)):
    if not files:
        raise HTTPException(status_code=400, detail="No files uploaded")
    job_id, d = _new_job_dir()
    paths: List[str] = []
    for f in files:
        p = _save_upload(f, d)
        img = Image.open(p)
        if img.mode in ("RGBA", "P"):
            img = img.convert("RGB")
            new_p = p.with_suffix(".jpg")
            img.save(new_p, "JPEG", quality=92)
            paths.append(str(new_p))
        else:
            paths.append(str(p))
    out = d / "combined.pdf"
    with out.open("wb") as fo:
        fo.write(img2pdf.convert(paths))
    return _build_response(out)


@api.post("/convert/merge-pdf", response_model=ConvertResponse)
async def merge_pdf(files: List[UploadFile] = File(...)):
    if len(files) < 2:
        raise HTTPException(status_code=400, detail="Upload at least two PDFs")
    job_id, d = _new_job_dir()
    writer = PdfWriter()
    for f in files:
        p = _save_upload(f, d)
        reader = PdfReader(str(p))
        for page in reader.pages:
            writer.add_page(page)
    out = d / "merged.pdf"
    with out.open("wb") as fo:
        writer.write(fo)
    return _build_response(out)


@api.post("/convert/split-pdf", response_model=ConvertResponse)
async def split_pdf(file: UploadFile = File(...)):
    job_id, d = _new_job_dir()
    src = _save_upload(file, d)
    reader = PdfReader(str(src))
    if len(reader.pages) <= 1:
        raise HTTPException(status_code=400, detail="PDF has only one page; nothing to split")
    out = d / f"{src.stem}_split.zip"
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf:
        for i, page in enumerate(reader.pages, 1):
            w = PdfWriter()
            w.add_page(page)
            buf = io.BytesIO()
            w.write(buf)
            zf.writestr(f"{src.stem}_page_{i}.pdf", buf.getvalue())
    return _build_response(out)


@api.post("/convert/compress-pdf", response_model=ConvertResponse)
async def compress_pdf(file: UploadFile = File(...)):
    job_id, d = _new_job_dir()
    src = _save_upload(file, d)
    out = d / f"{src.stem}_compressed.pdf"
    try:
        with pikepdf.open(str(src)) as pdf:
            pdf.save(str(out), compress_streams=True, object_stream_mode=pikepdf.ObjectStreamMode.generate)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Compression failed: {e}")
    return _build_response(out)


@api.post("/convert/pdf-to-excel", response_model=ConvertResponse)
async def pdf_to_excel(file: UploadFile = File(...)):
    job_id, d = _new_job_dir()
    src = _save_upload(file, d)
    out = d / f"{src.stem}.xlsx"
    wb = Workbook()
    wb.remove(wb.active)
    try:
        doc = fitz.open(str(src))
        for page_num, page in enumerate(doc, 1):
            ws = wb.create_sheet(title=f"Page {page_num}")
            tables = []
            try:
                tabs = page.find_tables()
                tables = list(tabs.tables) if tabs and tabs.tables else []
            except Exception:
                tables = []
            if tables:
                row_offset = 1
                for tbl in tables:
                    data = tbl.extract()
                    for r, row in enumerate(data, start=row_offset):
                        for c, val in enumerate(row, start=1):
                            ws.cell(row=r, column=c, value=val)
                    row_offset += len(data) + 2
            else:
                text = page.get_text("text")
                for r, line in enumerate(text.splitlines(), 1):
                    ws.cell(row=r, column=1, value=line)
        doc.close()
        if not wb.sheetnames:
            wb.create_sheet("Sheet1")
        wb.save(str(out))
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Conversion failed: {e}")
    return _build_response(out)


@api.post("/convert/excel-to-pdf", response_model=ConvertResponse)
async def excel_to_pdf(file: UploadFile = File(...)):
    job_id, d = _new_job_dir()
    src = _save_upload(file, d)
    out = _run_libreoffice(src, d, "pdf")
    return _build_response(out)


@api.post("/convert/pdf-to-powerpoint", response_model=ConvertResponse)
async def pdf_to_powerpoint(file: UploadFile = File(...)):
    job_id, d = _new_job_dir()
    src = _save_upload(file, d)
    out = d / f"{src.stem}.pptx"
    try:
        doc = fitz.open(str(src))
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        blank = prs.slide_layouts[6]
        img_dir = d / "slides"
        img_dir.mkdir(exist_ok=True)
        for i, page in enumerate(doc, 1):
            pix = page.get_pixmap(dpi=150)
            img_path = img_dir / f"slide_{i}.png"
            pix.save(str(img_path))
            slide = prs.slides.add_slide(blank)
            slide.shapes.add_picture(str(img_path), 0, 0, width=prs.slide_width, height=prs.slide_height)
        doc.close()
        prs.save(str(out))
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Conversion failed: {e}")
    return _build_response(out)


@api.post("/convert/powerpoint-to-pdf", response_model=ConvertResponse)
async def powerpoint_to_pdf(file: UploadFile = File(...)):
    job_id, d = _new_job_dir()
    src = _save_upload(file, d)
    out = _run_libreoffice(src, d, "pdf")
    return _build_response(out)


@api.post("/convert/image-converter", response_model=ConvertResponse)
async def image_converter(file: UploadFile = File(...), target_format: str = Form("png")):
    fmt = target_format.lower().strip()
    if fmt not in {"png", "jpg", "jpeg", "webp", "bmp"}:
        raise HTTPException(status_code=400, detail="Unsupported target format")
    pil_fmt = {"jpg": "JPEG", "jpeg": "JPEG", "png": "PNG", "webp": "WEBP", "bmp": "BMP"}[fmt]
    ext = "jpg" if fmt in ("jpg", "jpeg") else fmt
    job_id, d = _new_job_dir()
    src = _save_upload(file, d)
    img = Image.open(src)
    if pil_fmt == "JPEG" and img.mode in ("RGBA", "P", "LA"):
        img = img.convert("RGB")
    out = d / f"{src.stem}.{ext}"
    img.save(out, pil_fmt)
    return _build_response(out)


@api.get("/download/{file_id}")
async def download(file_id: str, background_tasks: BackgroundTasks):
    if not file_id.isalnum():
        raise HTTPException(status_code=400, detail="Invalid file id")
    job_dir = STORAGE_DIR / file_id
    if not job_dir.is_dir():
        raise HTTPException(status_code=404, detail="File expired or not found")
    candidates = sorted(
        [p for p in job_dir.iterdir() if p.is_file() and not p.name.startswith(".")],
        key=lambda p: p.stat().st_mtime, reverse=True,
    )
    if not candidates:
        raise HTTPException(status_code=404, detail="File expired or not found")
    target = candidates[0]
    return FileResponse(path=str(target), filename=target.name, media_type="application/octet-stream")


@api.post("/cleanup")
async def cleanup_now():
    _cleanup_old_jobs()
    return {"status": "ok"}


app.include_router(api)
app.add_middleware(
    CORSMiddleware,
    allow_origins=os.environ.get("CORS_ORIGINS", "*").split(","),
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.on_event("startup")
async def startup_event():
    _cleanup_old_jobs()
    logger.info("FlexiDoc storage at %s", STORAGE_DIR)


@app.on_event("shutdown")
async def shutdown_event():
    mongo_client.close()
